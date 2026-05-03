import os
from docx import Document

def export_docx(chapters, output_path="output/study_guide.docx"):
    doc = Document()
    doc.add_heading("Study Guide", 0)
    
    for chapter in chapters:
        doc.add_heading(chapter["title"], level=1)
        
        for line in chapter["content"]:
            # Basic support for bold could be added, but keeping it simple
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            else:
                # Remove ** for docx simplicity unless we parse it properly
                clean_line = line.replace("**", "")
                doc.add_paragraph(clean_line)
                
        if chapter.get("remember"):
            doc.add_heading("Remember this", level=2)
            for tip in chapter["remember"]:
                doc.add_paragraph(tip, style='List Bullet')
                
        if chapter.get("mnemonics"):
            doc.add_heading("Mnemonics", level=2)
            for keyword, mnemonic in chapter["mnemonics"].items():
                p = doc.add_paragraph()
                p.add_run(f"{keyword}: ").bold = True
                p.add_run(mnemonic)
                
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)

def export_md(chapters, output_path="output/study_guide.md"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("# Study Guide\n\n")
        
        for chapter in chapters:
            f.write(f"## {chapter['title']}\n\n")
            
            for line in chapter["content"]:
                f.write(f"{line}\n")
            f.write("\n")
            
            if chapter.get("remember"):
                f.write("### Remember this\n")
                for tip in chapter["remember"]:
                    f.write(f"- {tip}\n")
                f.write("\n")
                
            if chapter.get("mnemonics"):
                f.write("### Mnemonics\n")
                for keyword, mnemonic in chapter["mnemonics"].items():
                    f.write(f"- **{keyword}**: {mnemonic}\n")
                f.write("\n")

def export_notecards(chapters, output_path="output/notecards.txt"):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        for chapter in chapters:
            if chapter.get("mnemonics"):
                for keyword, mnemonic in chapter["mnemonics"].items():
                    # Replace newlines in mnemonic to keep it one line per card
                    flat_mnemonic = mnemonic.replace("\n", " ")
                    f.write(f"{keyword} | {flat_mnemonic}\n")
